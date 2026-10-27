"""Import an airport inventory Excel safely, preserving publisher audits and scene duplicates.

Put the newest .xlsx in data/airports/. The workbook must contain columns for
airport code, scene name or package folder, and installed version.
"""
from __future__ import annotations
import json, pathlib, re, sys, unicodedata
ROOT=pathlib.Path(__file__).resolve().parents[1]
DIR=ROOT/"data"/"airports"
OUT=ROOT/"public"/"airports.json"
ALIASES={
 "icao":{"icao","code","code icao","airport code","airport","code aeroport","aeroport"},
 "name":{"name","nom","scene","scene name","nom scene","airport name","nom aeroport"},
 "folder":{"folder","dossier","package","package name","nom dossier","community folder","repertoire"},
 "installed":{"installed","version","version installee","installed version","version locale","local version"},
 "publisher":{"publisher","editeur","developer","developpeur","auteur"},
 "region":{"region","continent","zone","category","categorie"},
}
CORRECTIONS={"BCGB":"LXGB","KABK":"KABQ"}
def norm(value):
 s=unicodedata.normalize("NFKD",str(value or "")).encode("ascii","ignore").decode().lower()
 return re.sub(r"[^a-z0-9]+"," ",s).strip()
def colmap(row):
 labels=[norm(x) for x in row]
 return {key:next((i for i,label in enumerate(labels) if label in names),None) for key,names in ALIASES.items()}
def is_sub_scene(*values):
 return any(re.search(r"\\(\\s*sub\\b",str(v or ""),re.I) for v in values)

def scene_type(code,name,folder):
 text=norm(" ".join((str(name or ""),str(folder or ""))))
 return "heliport" if re.search(r"\\b(heliport|heliports|helipad|helipads|heliportuaire|helisurface|helistop)\\b",text) or str(code).upper()=="58CA" else "airport"

def section_kind(value):
 label=norm(value)
 label=re.sub(r"\\s+\\d+\\s+\\d+$","",label)
 if label in ("aeroport","aeroports","airport","airports","airfields","aerodromes","aerodrome"):
  return "airports"
 if label in ("heliport","heliports","helipad","helipads","helistop","helisurfaces"):
  return "airports"
 if label in ("scenery","sceneries","decors","decor","cities","city","villes","ville","landmarks","other","autres","divers","world","landscape","landscapes"):
  return "other"
 return None

def parse_addons_linker(sheet):
 records=[]
 region=""
 country=""
 section=None
 rows=list(sheet.iter_rows(values_only=True))
 # Reconnaître les sections si l'export les fournit; ne jamais importer SUB.
 has_sections=any(section_kind(v) for row in rows for v in row[:5] if v is not None)
 for row in rows:
  if len(row)<13: continue
  heading=next((section_kind(v) for v in row[:5] if section_kind(v)),None)
  if heading and not row[6]: section=heading
  if has_sections and section!="airports": continue
  if row[3] and not row[4] and not row[5]:
   region=str(row[3]).split(" (")[0].strip()
  if row[4] and not row[5]:
   country=str(row[4]).split(" (")[0].strip()
  folder=str(row[5] or row[4] or "").strip()
  code=str(row[6] or "").strip().upper()
  version=str(row[9] or "").strip()
  name=str(row[11] or folder).strip()
  if is_sub_scene(code,name,folder): continue
  if not code or not version or not folder or not name: continue
  if not re.fullmatch(r"[A-Z0-9]{4}",code): continue
  original=code
  code=CORRECTIONS.get(code,code)
  records.append({"icao":code,"original":original,"name":name,"folder":folder,
                  "installed":version,"publisher":str(row[12] or "").strip(),
                  "region":region or country,"scene_type":scene_type(code,name,folder)})
 return records

def parse(book):
 from openpyxl import load_workbook
 wb=load_workbook(book,read_only=True,data_only=True)
 if any(str(cell.value or "").strip()=="ICAO Duplicate" for cell in wb.active[1]):
  return parse_addons_linker(wb.active)
 records=[]
 for sheet in wb:
  rows=sheet.iter_rows(values_only=True)
  for header in rows:
   if not any(v is not None for v in header): continue
   cols=colmap(header)
   if cols["installed"] is not None and (cols["folder"] is not None or cols["name"] is not None) and cols["icao"] is not None:
    break
  else: continue
  for row in rows:
   def get(key):
    idx=cols[key]
    return str(row[idx]).strip() if idx is not None and idx<len(row) and row[idx] is not None else ""
   code=get("icao").upper()
   folder=get("folder")
   name=get("name")
   version=get("installed")
   if is_sub_scene(code,name,folder): continue
   if not code and not folder and not name: continue
   if not code or not version or not (folder or name):
    raise ValueError(f"{sheet.title}: ligne incomplète : ICAO, scène/dossier et version requis")
   if not re.fullmatch(r"[A-Z0-9]{4}",code): continue
   original=code
   code=CORRECTIONS.get(code,code)
   records.append({"icao":code,"original":original,"name":name or folder,"folder":folder,"installed":version,"publisher":get("publisher"),"region":get("region"),"scene_type":scene_type(code,name,folder)})
 return records
def merge(records,old):
 by_folder={norm(a.get("folder")):a for a in old if a.get("folder")}
 by_identity={(a.get("icao"),norm(a.get("name")),norm(a.get("publisher"))):a for a in old}
 output=[]
 seen=set()
 for rec in records:
  key=norm(rec["folder"]) if rec["folder"] else (rec["icao"],norm(rec["name"]),norm(rec["publisher"]))
  if key in seen: raise ValueError(f"Doublon exact de scène dans Excel : {rec['name']}")
  seen.add(key)
  previous=by_folder.get(norm(rec["folder"])) if rec["folder"] else None
  if previous is None: previous=by_identity.get((rec["icao"],norm(rec["name"]),norm(rec["publisher"])))
  entry=dict(previous or {})
  entry.update({k:v for k,v in rec.items() if v or k in ("icao","original","installed")})
  if previous and previous.get("installed")!=entry["installed"]:
   for field in ("version_checked_at","version_comparison","version_check_state","server_version"):
    entry.pop(field,None)
   entry["status"]="UNK-V"
  entry["scene_type"]=rec.get("scene_type","airport")
  entry.setdefault("status","UNK-V")
  entry.setdefault("country",entry.get("region",""))
  output.append(entry)
 return output
def main():
 files=sorted(DIR.glob("*.xlsx"),key=lambda p:(p.stat().st_mtime,p.name),reverse=True) if DIR.exists() else []
 if not files:
  print("Aucun fichier Excel : inventaire actuel conservé.")
  return
 if len(files)>1: raise ValueError("Plusieurs fichiers Excel présents : conserver un seul fichier dans data/airports/")
 records=parse(files[0])
 if len(records)<10: raise ValueError(f"Import annulé : seulement {len(records)} scènes (minimum 10)")
 old=json.loads(OUT.read_text(encoding="utf-8"))
 result=merge(records,old)
 OUT.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
 print(f"Excel importé : {files[0].name}, {len(result)} scènes, {len(set(a['icao'] for a in result))} codes")
if __name__=="__main__":
 try: main()
 except Exception as exc:
  print(f"ERREUR IMPORT EXCEL : {exc}",file=sys.stderr)
  raise
